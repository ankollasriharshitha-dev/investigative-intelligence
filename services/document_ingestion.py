"""MVP document extraction and review contracts for synthetic/anonymized documents."""

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from domain.models import ValidationResult
from providers.synthetic import CSV_SCHEMAS, JSON_SCHEMAS


SUPPORTED_DOCUMENTS = {".pdf": "PDF", ".docx": "DOCX", ".txt": "TXT"}
ID_PREFIXES = ("P", "PH", "LOC", "V", "ORG", "ACC", "CI-")


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    document_type: str
    text: str
    entities: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    relationships: list[dict[str, Any]] = field(default_factory=list)
    warnings: tuple[str, ...] = ()
    validation: ValidationResult = ValidationResult(valid=False)


class DocumentIngestionService:
    """Extract explicit structured facts; approval remains a separate action."""

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_DOCUMENTS:
            return ExtractedDocument(filename, "Unsupported", "", warnings=("Unsupported document type. Use PDF, DOCX, or TXT.",))
        try:
            text, warnings = self._extract_text(suffix, content)
        except Exception as error:
            return ExtractedDocument(filename, SUPPORTED_DOCUMENTS[suffix], "", warnings=(f"Document extraction failed: {error}",))
        if not text.strip():
            warnings = warnings + ("No extractable text was found. OCR support is planned for a future production version.",)
        entities, relationships, parse_warnings = _parse_structured_facts(text, filename)
        warnings = warnings + tuple(parse_warnings)
        validation = _validate_extracted(entities, relationships)
        return ExtractedDocument(filename, SUPPORTED_DOCUMENTS[suffix], text, entities, relationships, warnings, validation)

    def package_files(self, extraction: ExtractedDocument) -> dict[str, bytes]:
        """Convert approved extracted facts into the existing canonical upload package."""
        files: dict[str, bytes] = {}
        for filename, columns in CSV_SCHEMAS.items():
            table_name = Path(filename).stem
            rows = extraction.entities.get(table_name, [])
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=list(columns) + ["source_type", "source_file", "source_case"], extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                enriched = dict(row)
                enriched.update({"source_type": "document", "source_file": extraction.filename, "source_case": _case_id(extraction)})
                writer.writerow(enriched)
            files[filename] = output.getvalue().encode("utf-8")
        cases = extraction.entities.get("cases", [])
        files["cases.json"] = json.dumps(cases, indent=2).encode("utf-8")
        return files

    @staticmethod
    def _extract_text(suffix: str, content: bytes) -> tuple[str, tuple[str, ...]]:
        if suffix == ".txt":
            return content.decode("utf-8"), ()
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            warnings = () if text.strip() else ("Text could not be extracted from this PDF. OCR support is planned for a future production version.",)
            return text, warnings
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + tables), ()


def _parse_structured_facts(text: str, filename: str) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]], list[str]]:
    entities: dict[str, list[dict[str, Any]]] = {Path(filename).stem: []}
    for table in (*[Path(name).stem for name in CSV_SCHEMAS], "cases"):
        entities.setdefault(table, [])
    relationships: list[dict[str, Any]] = []
    warnings: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or line.startswith("SYNTHETIC TRAINING DOCUMENT"):
            continue
        parts = [part.strip() for part in line.split("|")]
        kind = parts[0].upper().rstrip(":")
        if kind == "CASE" and len(parts) >= 3:
            entities["cases"].append({"case_id": parts[1], "title": parts[2], "status": parts[3] if len(parts) > 3 else "Open", "priority": parts[4] if len(parts) > 4 else "Medium", "opened_date": parts[5] if len(parts) > 5 else "2026-01-01"})
        elif kind in {"PERSON", "PHONE", "LOCATION", "VEHICLE", "ORGANIZATION", "ACCOUNT"}:
            record = _entity_record(kind, parts)
            if record:
                entities[_table_for(kind)].append(record)
        elif kind in {"CALL", "MESSAGE", "TRANSACTION", "EVENT"}:
            record = _relationship_record(kind, parts)
            if record:
                relationships.append(record)
                table = {"CALL": "calls", "MESSAGE": "messages", "TRANSACTION": "transactions", "EVENT": "events"}[kind]
                entities[table].append(record)
        else:
            explicit_ids = [token for token in re.findall(r"\b(?:P|PH|LOC|V|ORG|ACC)\d+\b|\bCI-\d{4}-\d{4}\b", line)]
            if explicit_ids:
                warnings.append(f"Unstructured line retained for review; no record generated: {line[:120]}")
    return entities, relationships, warnings


def _entity_record(kind: str, parts: list[str]) -> dict[str, Any] | None:
    case_id = parts[-1] if parts and parts[-1].startswith("CI-") else "CI-2026-0047"
    if kind == "PERSON" and len(parts) >= 3:
        return {"person_id": parts[1], "name": parts[2], "role": parts[3] if len(parts) > 3 and not parts[3].startswith("CI-") else "Document reference", "case_ids": case_id}
    if kind == "PHONE" and len(parts) >= 4:
        return {"phone_id": parts[1], "phone_number": parts[2], "owner_person_id": parts[3], "case_id": case_id}
    if kind == "LOCATION" and len(parts) >= 3:
        return {"location_id": parts[1], "label": parts[2], "latitude": parts[3] if len(parts) > 3 else "0", "longitude": parts[4] if len(parts) > 4 else "0"}
    if kind == "VEHICLE" and len(parts) >= 4:
        return {"vehicle_id": parts[1], "registration_number": parts[2], "owner_person_id": parts[3], "case_id": case_id}
    if kind == "ORGANIZATION" and len(parts) >= 3:
        return {"organization_id": parts[1], "name": parts[2], "sector": parts[3] if len(parts) > 3 else "Document reference", "case_ids": case_id}
    if kind == "ACCOUNT" and len(parts) >= 4:
        return {"account_id": parts[1], "account_label": parts[2], "owner_person_id": parts[3], "organization_id": parts[4] if len(parts) > 4 and not parts[4].startswith("CI-") else "", "case_id": case_id}
    return None


def _relationship_record(kind: str, parts: list[str]) -> dict[str, Any] | None:
    case_id = parts[-1] if parts and parts[-1].startswith("CI-") else "CI-2026-0047"
    if kind == "CALL" and len(parts) >= 4:
        return {"call_id": parts[1], "caller_id": parts[2], "receiver_id": parts[3], "timestamp": parts[4] if len(parts) > 4 else "2026-01-01T00:00:00", "case_id": case_id, "source": "document", "confidence": parts[5] if len(parts) > 5 else "0.70"}
    if kind == "MESSAGE" and len(parts) >= 4:
        return {"message_id": parts[1], "sender_phone_id": parts[2], "receiver_phone_id": parts[3], "timestamp": parts[4] if len(parts) > 4 else "2026-01-01T00:00:00", "case_id": case_id, "source": "document", "confidence": parts[5] if len(parts) > 5 else "0.70"}
    if kind == "TRANSACTION" and len(parts) >= 5:
        return {"transaction_id": parts[1], "from_account": parts[2], "to_account": parts[3], "amount": parts[4], "timestamp": parts[5] if len(parts) > 5 else "2026-01-01T00:00:00", "case_id": case_id, "source": "document", "confidence": parts[6] if len(parts) > 6 else "0.70"}
    if kind == "EVENT" and len(parts) >= 4:
        return {"event_id": parts[1], "person_id": parts[2], "event_type": parts[3], "location_id": parts[4] if len(parts) > 4 else "", "vehicle_id": parts[5] if len(parts) > 5 else "", "organization_id": parts[6] if len(parts) > 6 else "", "timestamp": parts[7] if len(parts) > 7 else "2026-01-01T00:00:00", "case_id": case_id, "source": "document", "confidence": parts[8] if len(parts) > 8 else "0.70"}
    return None


def _validate_extracted(entities: dict[str, list[dict[str, Any]]], relationships: list[dict[str, Any]]) -> ValidationResult:
    ids = {key: {row.get(column) for row in rows if row.get(column)} for key, rows in entities.items() for column in [_id_column(key)]}
    errors = []
    for kind, rows in entities.items():
        column = _id_column(kind)
        seen = [row.get(column) for row in rows]
        if len(seen) != len(set(seen)):
            errors.append(f"Duplicate extracted {column} values require review")
    for record in relationships:
        references = [value for key, value in record.items() if key.endswith("_id") and key not in {"call_id", "message_id", "transaction_id", "event_id"} and value]
        if not all(any(value in known for known in ids.values()) for value in references):
            errors.append(f"Extracted relationship {record.get('event_id') or record.get('call_id') or record.get('message_id') or record.get('transaction_id')} has an unresolved entity reference")
    return ValidationResult(valid=not errors, errors=tuple(errors))


def _id_column(table: str) -> str:
    return {"persons": "person_id", "phones": "phone_id", "locations": "location_id", "vehicles": "vehicle_id", "organizations": "organization_id", "accounts": "account_id", "calls": "call_id", "messages": "message_id", "transactions": "transaction_id", "events": "event_id", "cases": "case_id"}.get(table, "id")


def _table_for(kind: str) -> str:
    return {"PERSON": "persons", "PHONE": "phones", "LOCATION": "locations", "VEHICLE": "vehicles", "ORGANIZATION": "organizations", "ACCOUNT": "accounts"}[kind]


def _case_id(extraction: ExtractedDocument) -> str:
    cases = extraction.entities.get("cases", [])
    return cases[0].get("case_id", "") if cases else ""