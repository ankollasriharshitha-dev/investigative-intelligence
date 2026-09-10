import shutil
import stat
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from config.settings import get_settings
from domain.models import DataSourceType
from graph.analytics import GraphAnalytics
from graph.builder import InvestigationGraphBuilder
from graph.investigative import InvestigativeAnalytics
from providers.synthetic import SyntheticDataProvider
from services.investigation import InvestigationService
from services.data_management import DataManagementService
from services.ingestion import IngestionService
from services.document_ingestion import DocumentIngestionService
from providers.uploaded import UploadedDataProvider


def test_synthetic_provider_validates_external_schema() -> None:
    provider = SyntheticDataProvider(get_settings().synthetic_data_dir)
    result = provider.validate()

    assert result.valid
    assert result.checked_files == (
        "accounts.csv",
        "calls.csv",
        "cases.json",
        "events.csv",
        "locations.csv",
        "messages.csv",
        "organizations.csv",
        "persons.csv",
        "phones.csv",
        "transactions.csv",
        "vehicles.csv",
    )
    assert provider.get_source_info().source_type is DataSourceType.SYNTHETIC


def test_snapshot_is_loaded_from_files_through_provider() -> None:
    snapshot = SyntheticDataProvider(get_settings().synthetic_data_dir).load_snapshot()

    assert set(snapshot.tables) == {
        "accounts",
        "persons",
        "phones",
        "calls",
        "messages",
        "transactions",
        "locations",
        "vehicles",
        "organizations",
        "events",
    }
    assert len(snapshot.documents["cases"]) == 3


def test_investigation_service_calculates_active_source_counts() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))

    assert service.dashboard_statistics() == {
        "cases": 3,
        "persons": 18,
        "phones": 10,
        "vehicles": 6,
        "locations": 6,
        "organizations": 4,
        "accounts": 10,
        "relationships": 60,
    }


def test_invalid_relationship_endpoint_is_reported(tmp_path) -> None:
    data_dir = tmp_path / "synthetic"
    shutil.copytree(get_settings().synthetic_data_dir, data_dir)
    calls_path = data_dir / "calls.csv"
    calls_path.write_text(calls_path.read_text(encoding="utf-8").replace("C001,PH001,PH002", "C001,PH999,PH002", 1), encoding="utf-8")

    result = SyntheticDataProvider(data_dir).validate()

    assert not result.valid
    assert any("calls.csv.caller_id references unknown IDs" in error for error in result.errors)


def test_graph_is_built_from_active_provider_snapshot() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))

    graph = service.build_active_graph()

    assert graph.number_of_nodes() == 57
    assert graph.number_of_edges() == 134
    assert graph.nodes["P001"]["entity_type"] == "PERSON"
    assert graph.nodes["CASE001"]["entity_type"] == "CASE"
    assert "CALLED" in graph["PH001"]["PH002"]["relationship_types"]
    assert graph["PH001"]["PH002"]["case_id"] == "CASE001"


def test_graph_analytics_calculates_paths_hops_components_and_cross_case() -> None:
    snapshot = SyntheticDataProvider(get_settings().synthetic_data_dir).load_snapshot()
    analytics = GraphAnalytics(InvestigationGraphBuilder().build(snapshot))

    rankings = analytics.centrality()
    assert len(rankings) == 57
    assert rankings[0]["degree"] >= rankings[-1]["degree"]
    assert analytics.shortest_path("P001", "P010")
    assert analytics.hop_connections("P001", 2)
    assert analytics.hop_connections("P001", 3)
    assert len(analytics.connected_components()) == 2
    assert len(analytics.cross_case_entities()) == 9
    assert analytics.leads()


def test_service_graph_and_dashboard_use_same_active_provider() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))

    assert service.network_statistics()["nodes"] == service.build_active_graph().number_of_nodes()
    assert service.active_source().name == "Synthetic Demo Data"


def test_investigative_analytics_calculates_significance_and_clusters() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))
    analytics = service.investigative_analytics()

    rows = analytics.centrality()
    assert len(rows) == 57
    assert {"degree", "degree_centrality", "betweenness", "pagerank", "network_significance", "cases"} <= rows[0].keys()
    assert all(row["network_significance"] >= 0 for row in rows)
    assert analytics.clusters()
    assert any(row["cross_cluster_connections"] >= 0 for row in analytics.clusters())


def test_communication_transaction_location_vehicle_analytics_are_active_source_backed() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))
    analytics = service.investigative_analytics()

    assert len(analytics.communication_activity()) == 10
    assert len(analytics.transaction_activity()) == 10
    assert analytics.location_associations()
    assert len(analytics.vehicle_associations()) == 6


def test_lead_engine_returns_explainable_prioritized_signals() -> None:
    service = InvestigationService(SyntheticDataProvider(get_settings().synthetic_data_dir))
    leads = service.investigative_analytics().leads()

    assert leads
    assert {lead.status for lead in leads} == {"Requires Verification"}
    assert {lead.priority for lead in leads} <= {"Low", "Medium", "High"}
    assert {lead.lead_type for lead in leads} & {"NETWORK_BRIDGE", "HIGH_NETWORK_SIGNIFICANCE"}
    assert all(lead.reason and lead.evidence and lead.metrics for lead in leads)


def _test_upload_files() -> dict[str, bytes]:
    return {path.name: path.read_bytes() for path in Path("data/test_upload").iterdir() if path.is_file()}


def test_upload_validation_rejects_missing_files_and_broken_references(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    manager = DataManagementService(settings)
    files = _test_upload_files()
    files.pop("phones.csv")

    result, _ = manager.validate_files(files)

    assert not result.valid
    assert any("Missing required file: phones.csv" in error for error in result.errors)


def test_upload_validation_rejects_malformed_json_and_duplicate_ids(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    manager = DataManagementService(settings)
    files = _test_upload_files()
    files["cases.json"] = b"{not-json}"
    result, _ = manager.validate_files(files)
    assert not result.valid
    assert any("cases.json is not valid JSON" in error for error in result.errors)

    files = _test_upload_files()
    files["persons.csv"] = files["persons.csv"].replace(b"P905,Test Echo", b"P904,Test Echo")
    result, _ = manager.validate_files(files)
    assert not result.valid
    assert any("persons.csv.person_id contains duplicate IDs" in error for error in result.errors)


def test_upload_source_switch_changes_counts_graph_and_returns_to_synthetic(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    manager = DataManagementService(settings)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    report = service.import_dataset(_test_upload_files())

    assert report.validation.valid
    assert service.active_source().source_type is DataSourceType.SYNTHETIC
    assert service.activate_uploaded_dataset()
    assert isinstance(service.provider, UploadedDataProvider)
    assert service.dashboard_statistics()["persons"] == 5
    assert service.build_active_graph().number_of_nodes() < 57
    assert service.investigative_analytics().leads()

    service.clear_uploaded_dataset()

    assert service.active_source().source_type is DataSourceType.SYNTHETIC
    assert service.dashboard_statistics()["persons"] == 18
    assert service.build_active_graph().number_of_nodes() == 57


def test_clear_uploaded_removes_nested_read_only_files_and_restores_synthetic(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    report = service.import_dataset(_test_upload_files())
    assert report.validation.valid
    assert service.activate_uploaded_dataset()

    nested = settings.uploads_data_dir / "current" / "nested" / "locked.txt"
    nested.parent.mkdir(parents=True)
    nested.write_text("synthetic uploaded fixture", encoding="utf-8")
    nested.chmod(stat.S_IREAD)

    service.clear_uploaded_dataset()

    assert not (settings.uploads_data_dir / "current").exists()
    assert service.active_source().source_type is DataSourceType.SYNTHETIC
    assert service.dashboard_statistics()["persons"] == 18


def test_upload_append_and_update_use_stable_ids(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    service.import_dataset(_test_upload_files())
    service.activate_uploaded_dataset()
    files = _test_upload_files()
    append = service.import_dataset(files, "append")
    assert append.records_added == 0
    assert append.duplicates_skipped > 0
    files["persons.csv"] = files["persons.csv"].replace(b"P901,Test Alpha", b"P901,Updated Alpha")

    update = service.import_dataset(files, "update")

    assert update.records_updated >= 5
    assert service.load_active_dataset().tables["persons"].loc[lambda frame: frame.person_id == "P901", "name"].iloc[0] == "Updated Alpha"
    service.clear_uploaded_dataset()


def test_append_preserves_existing_cases_and_graph_refreshes_after_new_record(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    service.import_dataset(_test_upload_files())
    service.activate_uploaded_dataset()
    files = _test_upload_files()
    files["persons.csv"] = files["persons.csv"].rstrip(b"\n") + b"\nP906,Test Foxtrot,Observer,CASE901\n"
    append = service.import_dataset(files, "append")

    assert append.records_added >= 1
    assert "CASE901" in service.load_active_dataset().documents["cases"][0]["case_id"]
    assert len(service.load_active_dataset().tables["persons"]) == 6
    assert service.build_active_graph().number_of_nodes() > 0
    service.clear_uploaded_dataset()


def test_empty_valid_dataset_builds_empty_graph_without_analytics_crash(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    manager = DataManagementService(settings)
    files = {name: content for name, content in _test_upload_files().items()}
    files["persons.csv"] = b"person_id,name,role,case_ids\n"
    files["phones.csv"] = b"phone_id,phone_number,owner_person_id,case_id\n"
    files["calls.csv"] = b"call_id,caller_id,receiver_id,timestamp,case_id,source,confidence\n"
    files["messages.csv"] = b"message_id,sender_phone_id,receiver_phone_id,timestamp,case_id,source,confidence\n"
    files["transactions.csv"] = b"transaction_id,from_account,to_account,amount,timestamp,case_id,source,confidence\n"
    files["locations.csv"] = b"location_id,label,latitude,longitude\n"
    files["vehicles.csv"] = b"vehicle_id,registration_number,owner_person_id,case_id\n"
    files["organizations.csv"] = b"organization_id,name,sector,case_ids\n"
    files["accounts.csv"] = b"account_id,account_label,owner_person_id,organization_id,case_id\n"
    files["events.csv"] = b"event_id,person_id,event_type,location_id,vehicle_id,organization_id,timestamp,case_id,source,confidence\n"
    files["cases.json"] = b"[]"
    report = manager.import_dataset(files)
    assert report.validation.valid
    assert manager.activate_uploaded()
    provider = UploadedDataProvider(manager.current_dir)
    assert provider.load_snapshot().documents["cases"] == []
    manager.clear_uploaded()


def test_document_ingestion_extracts_txt_docx_and_pdf() -> None:
    service = DocumentIngestionService()
    fixture = Path("data/test_documents/northbridge_college_ci_2026_0047.txt")

    text_result = service.extract(fixture.name, fixture.read_bytes())
    document = Document()
    document.add_paragraph(fixture.read_text(encoding="utf-8"))
    docx_buffer = BytesIO()
    document.save(docx_buffer)
    docx_result = service.extract("college_case.docx", docx_buffer.getvalue())
    pdf_buffer = BytesIO()
    PdfWriter().add_blank_page(width=612, height=792)
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(pdf_buffer)
    pdf_result = service.extract("empty.pdf", pdf_buffer.getvalue())

    assert text_result.document_type == "TXT"
    assert docx_result.document_type == "DOCX"
    assert len(text_result.entities["persons"]) == 6
    assert len(docx_result.relationships) == 18
    assert pdf_result.document_type == "PDF"
    assert any("OCR support" in warning for warning in pdf_result.warnings)


def test_document_ingestion_rejects_unsupported_and_empty_documents() -> None:
    service = DocumentIngestionService()

    unsupported = service.extract("malware.exe", b"not accepted")
    empty = service.extract("empty.txt", b"")

    assert unsupported.document_type == "Unsupported"
    assert unsupported.warnings
    assert any("No extractable text" in warning for warning in empty.warnings)


def test_document_approval_uses_uploaded_pipeline_and_refreshes_graph_and_leads(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    document_service = DocumentIngestionService()
    fixture = Path("data/test_documents/northbridge_college_ci_2026_0047.docx")
    extraction = document_service.extract(fixture.name, fixture.read_bytes())

    assert extraction.validation.valid
    report = service.approve_document(extraction, "append")

    assert report.validation.valid
    assert report.dataset_version == 1
    assert service.active_source().source_type is DataSourceType.SYNTHETIC
    assert service.activate_uploaded_dataset()
    assert service.active_source().source_type is DataSourceType.UPLOADED
    assert service.dashboard_statistics()["persons"] == 6
    assert service.build_active_graph().number_of_nodes() > 0
    assert service.investigative_analytics().leads()
    snapshot = service.load_active_dataset()
    assert snapshot.tables["persons"].iloc[0]["source_type"] == "document"
    assert snapshot.tables["persons"].iloc[0]["source_file"] == fixture.name
    service.clear_uploaded_dataset()


def test_dataset_versions_ingestion_records_and_daily_update_refresh(tmp_path) -> None:
    settings = settings_for_upload_test(tmp_path)
    service = InvestigationService(SyntheticDataProvider(settings.synthetic_data_dir), settings)
    manager = service.data_management
    report = service.import_dataset(_test_upload_files())
    assert report.ingestion_id and report.dataset_version == 1
    service.activate_uploaded_dataset()
    assert manager.dataset_metadata().version == 1
    initial_nodes = service.build_active_graph().number_of_nodes()

    daily = IngestionService(manager)
    update = daily.apply_daily_update(_test_upload_files(), "append")

    assert update.duplicates_skipped > 0
    assert manager.dataset_metadata().version == 2
    assert service.build_active_graph().number_of_nodes() == initial_nodes
    assert manager.ingestion_history()[-1]["operation"] == "APPEND"
    service.clear_uploaded_dataset()


def settings_for_upload_test(tmp_path):
    return type(get_settings())(tmp_path, Path("data/synthetic"), tmp_path / "uploads")
